"""On-demand exports from the canonical Plan, PlanDay, and PlanItem records."""

from io import BytesIO
from pathlib import Path

import arabic_reshaper
from bidi.algorithm import get_display
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


FONT_DIR = Path(__file__).resolve().parent / "fonts"
KIND_LABELS = {"STUDY": "مطالعه", "TEST": "تست", "REVIEW": "مرور", "EXAM": "آزمون", "EVENT": "رویداد"}
STATUS_LABELS = {"DRAFT": "پیش‌نویس", "PUBLISHED": "منتشرشده"}


def display_name(user):
    return user.get_full_name().strip() or user.username


def item_subject(item):
    return item.subject.name if item.subject_id else ""


def item_chapter(item):
    return item.chapter.name if item.chapter_id else ""


def item_topic(item):
    return item.topic.name if item.topic_id else ""


def item_details(item):
    parts = [item.title, item_subject(item), item_chapter(item), item_topic(item)]
    return "، ".join(part for part in parts if part)


def item_metrics(item):
    parts = []
    if item.planned_duration_minutes is not None:
        parts.append(f"{item.planned_duration_minutes} دقیقه")
    if item.test_count is not None:
        parts.append(f"{item.test_count} تست")
    if item.start_time and item.end_time:
        parts.append(f"{item.start_time:%H:%M} تا {item.end_time:%H:%M}")
    return " | ".join(parts)


def date_label(value):
    return f"{value.day:02d} ماه {value.month:02d} سال {value.year}"


def _visual(text):
    return get_display(arabic_reshaper.reshape(str(text)), base_dir="R")


def _register_fonts():
    if "NotoNaskh" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("NotoNaskh", str(FONT_DIR / "NotoNaskhArabic-Regular.ttf")))
        pdfmetrics.registerFont(TTFont("NotoNaskhBold", str(FONT_DIR / "NotoNaskhArabic-Bold.ttf")))


def render_pdf(plan):
    _register_fonts()
    output = BytesIO()
    pdf = canvas.Canvas(output, pagesize=A4, pageCompression=1)
    pdf.setTitle(f"Plan {plan.pk}")
    page_width, page_height = A4
    right = page_width - 45
    left = 45
    available = right - left
    y = page_height - 52

    def next_page_if_needed(height=22):
        nonlocal y
        if y - height < 52:
            pdf.showPage()
            y = page_height - 52

    def line(text, size=11, bold=False, gap=19):
        nonlocal y
        font = "NotoNaskhBold" if bold else "NotoNaskh"
        pdf.setFont(font, size)
        words = str(text).replace("\r", "").split()
        if not words:
            y -= gap
            return
        current = ""
        for word in words:
            candidate = f"{current} {word}".strip()
            if current and pdfmetrics.stringWidth(_visual(candidate), font, size) > available:
                next_page_if_needed(gap)
                pdf.drawRightString(right, y, _visual(current))
                y -= gap
                current = word
            else:
                current = candidate
        if current:
            next_page_if_needed(gap)
            pdf.drawRightString(right, y, _visual(current))
            y -= gap

    line("برنامه هفتگی دانش‌آموز", size=17, bold=True, gap=28)
    line(f"دانش‌آموز: {display_name(plan.student.user)}", bold=True)
    line(f"مشاور: {display_name(plan.counselor.user)}")
    line(f"بازه برنامه: از {date_label(plan.start_date)} تا {date_label(plan.end_date)}")
    line(f"وضعیت: {STATUS_LABELS[plan.status]}")
    y -= 8
    for day in plan.days.all():
        next_page_if_needed(50)
        pdf.setStrokeColor(colors.HexColor("#AAB5BB"))
        pdf.line(left, y + 10, right, y + 10)
        y -= 12
        line(f"روز {date_label(day.date)}", size=13, bold=True, gap=24)
        for index, item in enumerate(day.items.all(), 1):
            line(f"{index}. {KIND_LABELS[item.kind]} — {item_details(item)}", bold=True)
            metrics = item_metrics(item)
            if metrics:
                line(metrics, size=10, gap=17)
            if item.note:
                for paragraph in item.note.splitlines():
                    if paragraph.strip():
                        line(f"یادداشت: {paragraph}", size=10, gap=17)
            y -= 4
        y -= 8
    pdf.save()
    return output.getvalue()


def _safe_text(value):
    """Keep user-entered spreadsheet text from being interpreted as a formula."""
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def render_excel(plan):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "برنامه"
    sheet.sheet_view.rightToLeft = True
    sheet.append(["دانش‌آموز", _safe_text(display_name(plan.student.user))])
    sheet.append(["مشاور", _safe_text(display_name(plan.counselor.user))])
    sheet.append(["شروع", plan.start_date, "پایان", plan.end_date])
    sheet.append(["وضعیت", STATUS_LABELS[plan.status]])
    sheet.append([])
    headers = ("تاریخ", "نوع فعالیت", "عنوان", "درس", "فصل", "مبحث", "مدت (دقیقه)", "تعداد تست", "شروع", "پایان", "یادداشت")
    sheet.append(headers)
    for day in plan.days.all():
        for item in day.items.all():
            sheet.append([
                day.date, KIND_LABELS[item.kind], _safe_text(item.title), _safe_text(item_subject(item)),
                _safe_text(item_chapter(item)), _safe_text(item_topic(item)), item.planned_duration_minutes,
                item.test_count, item.start_time, item.end_time, _safe_text(item.note),
            ])
    sheet.freeze_panes = "A7"
    sheet.auto_filter.ref = f"A6:K{max(sheet.max_row, 6)}"
    widths = (15, 16, 28, 20, 20, 20, 17, 14, 13, 13, 45)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cell in sheet[6]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="285B67")
    for row in sheet:
        for cell in row:
            cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
    for row in sheet.iter_rows(min_row=7):
        for index in (0,):
            row[index].number_format = "yyyy/mm/dd"
        for index in (8, 9):
            row[index].number_format = "hh:mm"
    sheet["B3"].number_format = "yyyy/mm/dd"
    sheet["D3"].number_format = "yyyy/mm/dd"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
