from datetime import date, time
from io import BytesIO

from django.contrib.auth import get_user_model
from django.db import connection
from django.test.utils import CaptureQueriesContext
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from apps.accounts.models import CounselorProfile, StudentProfile
from apps.academics.models import Field, Grade, Subject
from apps.planning.exports import item_details, item_metrics
from apps.planning.models import Plan, PlanDay, PlanItem


User = get_user_model()


class PlanExportTests(APITestCase):
    def setUp(self):
        self.admin = User.objects.create_user("export_admin", role=User.Role.ADMIN)
        self.counselor_user = User.objects.create_user("export_counselor", first_name="علی", last_name="رضایی", role=User.Role.COUNSELOR)
        self.counselor = CounselorProfile.objects.create(user=self.counselor_user)
        self.other_counselor_user = User.objects.create_user("export_other_counselor", role=User.Role.COUNSELOR)
        CounselorProfile.objects.create(user=self.other_counselor_user)
        self.student_user = User.objects.create_user("export_student", first_name="سارا", last_name="احمدی")
        self.student = StudentProfile.objects.create(user=self.student_user, counselor=self.counselor)
        self.other_student_user = User.objects.create_user("export_other_student")
        StudentProfile.objects.create(user=self.other_student_user, counselor=self.counselor)
        grade = Grade.objects.create(name="11")
        field = Field.objects.create(grade=grade, name="ریاضی")
        self.subject = Subject.objects.create(field=field, name="فیزیک")
        self.plan = Plan.objects.create(student=self.student, counselor=self.counselor, start_date=date(2026, 10, 5), end_date=date(2026, 10, 11))
        self.day = PlanDay.objects.create(plan=self.plan, date=date(2026, 10, 5))
        self.item = PlanItem.objects.create(
            plan_day=self.day, kind=PlanItem.Kind.TEST, subject=self.subject,
            planned_duration_minutes=60, test_count=20, ordering=0, note="مرور فرمول‌ها",
            start_time=time(15, 30), end_time=time(16, 30),
        )
        PlanItem.objects.create(plan_day=self.day, kind=PlanItem.Kind.EVENT, title="باشگاه", ordering=1)

    def url(self, kind, plan=None):
        return f"/api/v1/planning/plans/{(plan or self.plan).pk}/export/{kind}/"

    def test_authorized_counselor_pdf_contains_embedded_font_and_plan_data(self):
        self.client.force_authenticate(self.counselor_user)
        response = self.client.get(self.url("pdf"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertIn('attachment; filename="plan-', response["Content-Disposition"])
        self.assertTrue(response.content.startswith(b"%PDF-"))
        self.assertIn(b"NotoNaskh", response.content)
        self.assertGreater(len(response.content), 10000)
        self.assertEqual(item_details(self.item), "فیزیک")
        self.assertIn("20 تست", item_metrics(self.item))

    def test_authorized_counselor_excel_has_canonical_rows(self):
        self.client.force_authenticate(self.counselor_user)
        response = self.client.get(self.url("excel"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertTrue(response.content.startswith(b"PK"))
        sheet = load_workbook(BytesIO(response.content), data_only=True).active
        self.assertTrue(sheet.sheet_view.rightToLeft)
        self.assertEqual(sheet["B1"].value, "سارا احمدی")
        self.assertEqual(sheet["B2"].value, "علی رضایی")
        self.assertEqual(sheet["D7"].value, "فیزیک")
        self.assertEqual(sheet["G7"].value, 60)
        self.assertEqual(sheet["H7"].value, 20)
        self.assertEqual(sheet["K7"].value, "مرور فرمول‌ها")
        self.assertEqual(sheet["C8"].value, "باشگاه")
        self.assertIsNone(sheet["G8"].value)

    def test_export_permission_boundaries(self):
        for kind in ("pdf", "excel"):
            with self.subTest(kind=kind):
                self.plan.status = Plan.Status.DRAFT
                self.plan.published_at = None
                self.plan.save()
                self.client.force_authenticate(user=None)
                self.assertEqual(self.client.get(self.url(kind)).status_code, 401)
                self.client.force_authenticate(self.other_counselor_user)
                self.assertEqual(self.client.get(self.url(kind)).status_code, 404)
                self.client.force_authenticate(self.student_user)
                self.assertEqual(self.client.get(self.url(kind)).status_code, 404)
                self.plan.status = Plan.Status.PUBLISHED
                self.plan.published_at = timezone.now()
                self.plan.save()
                self.assertEqual(self.client.get(self.url(kind)).status_code, 200)
                self.client.force_authenticate(self.other_student_user)
                self.assertEqual(self.client.get(self.url(kind)).status_code, 404)
                self.client.force_authenticate(self.admin)
                self.assertEqual(self.client.get(self.url(kind)).status_code, 200)

    def test_export_queries_stay_bounded_as_items_grow(self):
        for index in range(12):
            PlanItem.objects.create(plan_day=self.day, kind=PlanItem.Kind.STUDY, subject=self.subject, planned_duration_minutes=45, ordering=index + 2)
        self.client.force_authenticate(self.counselor_user)
        for kind in ("pdf", "excel"):
            with self.subTest(kind=kind), CaptureQueriesContext(connection) as queries:
                response = self.client.get(self.url(kind))
            self.assertEqual(response.status_code, 200)
            self.assertLessEqual(len(queries), 6)

    def test_excel_escapes_user_entered_formulas(self):
        self.item.note = "=SUM(1,1)"
        self.item.save()
        self.client.force_authenticate(self.counselor_user)
        response = self.client.get(self.url("excel"))
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet["K7"].data_type, "s")
        self.assertEqual(sheet["K7"].value, "'=SUM(1,1)")
